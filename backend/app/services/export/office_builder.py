"""Office exporters (spec §4.6). PptxGenJS-equivalent server path uses
python-pptx (pptx) and openpyxl (xlsx). Fully local, no conversion service.
"""

from __future__ import annotations

import io
from typing import Iterable

from pptx import Presentation
from pptx.util import Inches, Pt
from openpyxl import Workbook


def build_pptx(slides: Iterable[tuple[str, list[str]]]) -> bytes:
    prs = Presentation()
    for title, bullets in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        body = slide.placeholders[1]
        body.text = ""
        for b in bullets:
            p = body.text_frame.add_paragraph()
            p.text = b
            p.font.size = Pt(18)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_xlsx(rows: list[list[str]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
