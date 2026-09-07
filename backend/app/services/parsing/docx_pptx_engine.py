"""DOCX / PPTX text extraction (spec Section 4.1).

Uses python-docx and python-pptx. Both run fully locally; no API keys.
"""
from __future__ import annotations

from docx import Document as DocxDocument
from pptx import Presentation


def extract_docx_text(path) -> str:
    """Extract paragraph text from a .docx file (path or file-like)."""
    doc = DocxDocument(path)
    paras = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    return "\n".join(paras)


def extract_pptx_text(path) -> str:
    """Extract text frames from every slide of a .pptx file."""
    prs = Presentation(path)
    out: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs)
                if text and text.strip():
                    out.append(text.strip())
    return "\n".join(out)
